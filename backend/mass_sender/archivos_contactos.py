"""
Lectura de archivos de contactos para la importación.

Formatos aceptados:
  - Excel moderno: .xlsx .xlsm .xltx .xltm  (openpyxl)
  - Excel 97-2003: .xls                       (xlrd)
  - Excel binario: .xlsb                      (pyxlsb)
  - LibreOffice / OpenOffice: .ods            (lector propio, sin dependencias)
  - Texto: .csv .txt .tsv                     (separador detectado solo)

Los Excel pueden pesar hasta 1 GB (por ejemplo, con imágenes o formatos):
se leen desde el archivo temporal que Django crea durante la subida, por
partes, sin cargarlo entero en memoria. Django borra ese temporal al terminar
la petición, así que no se guarda nada en el servidor. Lo que sí tiene tope es
la cantidad de filas (MAX_FILAS), porque todas viajan a la vista previa.
"""
import csv
import io
import re
import zipfile
from datetime import date, datetime
from xml.etree import ElementTree as ET

EXT_EXCEL = {'.xlsx', '.xlsm', '.xltx', '.xltm'}
EXT_TEXTO = {'.csv', '.txt', '.tsv'}
EXT_HOJAS = EXT_EXCEL | {'.xls', '.xlsb', '.ods'}      # Excel y similares
EXTENSIONES = EXT_HOJAS | EXT_TEXTO

MAX_BYTES_EXCEL = 1024 * 1024 * 1024    # 1 GB
MAX_BYTES_TEXTO = 50 * 1024 * 1024      # 50 MB
MAX_FILAS = 200_000
MAX_COLUMNAS = 60


class ArchivoInvalido(Exception):
    pass


def _texto(valor):
    """Celda -> texto. Los números enteros pierden el '.0' (los teléfonos
    que Excel guarda como número llegan como 983706769.0)."""
    if valor is None:
        return ''
    if isinstance(valor, bool):
        return 'SI' if valor else 'NO'
    if isinstance(valor, float):
        if valor.is_integer():
            return str(int(valor))
        return repr(valor)
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, datetime):
        return valor.strftime('%d/%m/%Y %H:%M') if (valor.hour or valor.minute) else valor.strftime('%d/%m/%Y')
    if isinstance(valor, date):
        return valor.strftime('%d/%m/%Y')
    return str(valor).strip()


def _limpiar(filas):
    """Quita filas vacías, recorta columnas vacías al final y separa el encabezado."""
    filas = [[_texto(c) for c in f[:MAX_COLUMNAS]] for f in filas]
    filas = [f for f in filas if any(c for c in f)]
    if not filas:
        raise ArchivoInvalido('El archivo no tiene datos.')
    ancho = max(len(f) for f in filas)
    while ancho > 1 and not any(len(f) >= ancho and f[ancho - 1] for f in filas):
        ancho -= 1
    filas = [(f + [''] * ancho)[:ancho] for f in filas]
    encabezados = [h or f'Columna {i + 1}' for i, h in enumerate(filas[0])]
    return encabezados, filas[1:]


def _limitar(iterable):
    filas = []
    for fila in iterable:
        filas.append(list(fila))
        if len(filas) > MAX_FILAS + 1:
            raise ArchivoInvalido(f'El archivo tiene más de {MAX_FILAS:,} filas. Divídelo en partes.'.replace(',', '.'))
    return filas


# ── Lectores por formato ────────────────────────────────────────────────
def _leer_xlsx(ruta):
    from openpyxl import load_workbook
    # Se pasa el archivo abierto (no la ruta) para que openpyxl no rechace
    # un xlsx guardado con otra extensión (.xls, .upload, ...).
    with open(ruta, 'rb') as fh:
        try:
            wb = load_workbook(fh, read_only=True, data_only=True)
        except Exception:
            raise ArchivoInvalido('No se pudo abrir el Excel. ¿Está dañado o protegido con contraseña?')
        try:
            ws = wb.worksheets[0]
            return _limitar(ws.iter_rows(values_only=True))
        finally:
            wb.close()


def _leer_xls(ruta):
    import xlrd
    try:
        libro = xlrd.open_workbook(filename=ruta, on_demand=True)
    except Exception:
        raise ArchivoInvalido('No se pudo abrir el Excel (.xls). ¿Está dañado o protegido con contraseña?')
    hoja = libro.sheet_by_index(0)

    def fila(r):
        out = []
        for c in range(hoja.ncols):
            celda = hoja.cell(r, c)
            if celda.ctype == xlrd.XL_CELL_DATE:
                try:
                    out.append(xlrd.xldate.xldate_as_datetime(celda.value, libro.datemode))
                    continue
                except Exception:
                    pass
            out.append(celda.value)
        return out
    return _limitar(fila(r) for r in range(hoja.nrows))


def _leer_xlsb(ruta):
    from pyxlsb import open_workbook
    try:
        with open_workbook(ruta) as wb:
            with wb.get_sheet(1) as hoja:
                return _limitar([c.v for c in fila] for fila in hoja.rows())
    except ArchivoInvalido:
        raise
    except Exception:
        raise ArchivoInvalido('No se pudo abrir el Excel binario (.xlsb).')


def _leer_ods(ruta):
    """Lee la primera hoja de un .ods (un zip con content.xml), por partes."""
    T = '{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
    O = '{urn:oasis:names:tc:opendocument:xmlns:office:1.0}'
    P = '{urn:oasis:names:tc:opendocument:xmlns:text:1.0}p'
    try:
        z = zipfile.ZipFile(ruta)
        xml = z.open('content.xml')
    except Exception:
        raise ArchivoInvalido('No se pudo abrir el archivo .ods.')

    filas = []
    hojas = 0
    with z, xml:
        for evento, elem in ET.iterparse(xml, events=('start', 'end')):
            if evento == 'start':
                if elem.tag == f'{T}table':
                    hojas += 1
                continue
            if hojas != 1:
                if elem.tag == f'{T}table':
                    break            # solo la primera hoja
                continue
            if elem.tag != f'{T}table-row':
                continue
            rep_f = min(int(elem.get(f'{T}number-rows-repeated', '1')), 1000)
            celdas = []
            for celda in elem:
                if celda.tag not in (f'{T}table-cell', f'{T}covered-table-cell'):
                    continue
                rep_c = min(int(celda.get(f'{T}number-columns-repeated', '1')), MAX_COLUMNAS)
                if celda.get(f'{O}value-type') in ('float', 'percentage', 'currency'):
                    try:
                        valor = float(celda.get(f'{O}value'))
                    except (TypeError, ValueError):
                        valor = ''
                else:
                    valor = '\n'.join(''.join(p.itertext()) for p in celda.findall(P))
                celdas.extend([valor] * rep_c)
            elem.clear()
            if any(_texto(c) for c in celdas):
                filas.extend([celdas] * rep_f)
            if len(filas) > MAX_FILAS + 1:
                raise ArchivoInvalido(f'El archivo tiene más de {MAX_FILAS:,} filas. Divídelo en partes.'.replace(',', '.'))
    if hojas == 0:
        raise ArchivoInvalido('El archivo .ods no tiene hojas.')
    return filas


def _leer_texto(datos, separador=''):
    for codif in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            texto = datos.decode(codif)
            break
        except UnicodeDecodeError:
            continue
    if not separador:
        muestra = texto[:20000]
        try:
            separador = csv.Sniffer().sniff(muestra, delimiters=',;\t|').delimiter
        except csv.Error:
            # Se elige el separador que más aparece en la primera línea
            primera = muestra.splitlines()[0] if muestra else ''
            separador = max(',;\t|', key=primera.count) if primera else ','
    lector = csv.reader(io.StringIO(texto, newline=''), delimiter=separador)
    return _limitar(lector)


def leer_archivo(archivo, separador=''):
    """Devuelve (encabezados, filas) a partir de un UploadedFile de Django."""
    nombre = (archivo.name or '').lower()
    ext = '.' + nombre.rsplit('.', 1)[-1] if '.' in nombre else ''
    if ext not in EXTENSIONES:
        raise ArchivoInvalido('Formato no soportado. Usa Excel (.xlsx, .xls, .xlsm, .xlsb), .ods o .csv.')
    if ext in EXT_HOJAS and archivo.size > MAX_BYTES_EXCEL:
        raise ArchivoInvalido('El Excel supera el límite de 1 GB.')
    if ext in EXT_TEXTO and archivo.size > MAX_BYTES_TEXTO:
        raise ArchivoInvalido('El CSV supera el límite de 50 MB. Guárdalo como Excel (.xlsx) para subir hasta 1 GB.')

    if ext in EXT_TEXTO:
        return _limpiar(_leer_texto(archivo.read(), separador))

    # Hojas de cálculo: se leen desde una ruta en disco (el temporal de la
    # subida). Si el archivo es pequeño y Django lo tuvo en memoria, se
    # escribe un temporal propio que se borra al terminar.
    import os
    import tempfile
    propio = None
    if hasattr(archivo, 'temporary_file_path'):
        ruta = archivo.temporary_file_path()
    else:
        fd, propio = tempfile.mkstemp(suffix=ext)
        with os.fdopen(fd, 'wb') as out:
            for trozo in archivo.chunks():
                out.write(trozo)
        ruta = propio
    try:
        with open(ruta, 'rb') as fh:
            cabecera = fh.read(512)
        if ext in EXT_EXCEL:
            filas = _leer_xlsx(ruta)
        elif ext == '.xls':
            # Algunos sistemas exportan HTML, CSV o xlsx con extensión .xls
            if cabecera[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                filas = _leer_xls(ruta)
            elif cabecera[:2] == b'PK':
                filas = _leer_xlsx(ruta)
            elif re.match(rb'\s*<', cabecera[:200]):
                raise ArchivoInvalido('Este .xls es en realidad una página web. Ábrelo en Excel y guárdalo como .xlsx.')
            else:
                if os.path.getsize(ruta) > MAX_BYTES_TEXTO:
                    raise ArchivoInvalido('Este .xls es en realidad texto y supera 50 MB. Guárdalo como .xlsx.')
                with open(ruta, 'rb') as fh:
                    filas = _leer_texto(fh.read(), separador or '\t')
        elif ext == '.xlsb':
            filas = _leer_xlsb(ruta)
        else:
            filas = _leer_ods(ruta)
    finally:
        if propio:
            try:
                os.remove(propio)
            except OSError:
                pass
    return _limpiar(filas)
