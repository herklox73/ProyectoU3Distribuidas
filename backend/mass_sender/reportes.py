"""
Reportes de mensajería: estadísticas y descargas en Excel y PDF.

Los archivos se generan EN MEMORIA en el momento de la descarga y se envían
directo al navegador. No se guarda nada en disco ni en la base de datos: los
datos ya existen en la tabla de mensajes, así que cada descarga se arma de
nuevo con la información al día y el servidor no acumula archivos.
"""
import io
from datetime import datetime, timedelta

from django.db.models import Count, Q
from django.db.models.functions import ExtractHour, TruncDate
from django.utils import timezone

from .models import Contact, Message
from .telefonos import normalizar_o_igual

# Estados de un mensaje saliente, en orden de avance
ESTADOS = [
    ('pending', 'Pendiente'),
    ('sent', 'Enviado'),
    ('delivered', 'Entregado'),
    ('read', 'Leído'),
    ('failed', 'Fallido'),
]
ETIQUETA = dict(ESTADOS)

# Colores (mismos que la página): rampa azul de menos a más avance,
# gris para pendiente y rojo de alerta para fallido.
COLOR = {
    'pending': '#a3a29b',
    'sent': '#86b6ef',
    'delivered': '#2a78d6',
    'read': '#104281',
    'failed': '#d03b3b',
}

MAX_FILAS_EXCEL = 20000   # detalle de mensajes en el Excel
MAX_FILAS_PDF = 150       # el PDF es para leer: el detalle completo va en el Excel


def rango_fechas(request):
    """Lee ?desde=AAAA-MM-DD&hasta=AAAA-MM-DD (por defecto, los últimos 30 días)."""
    hoy = timezone.localdate()
    try:
        desde = datetime.strptime(request.GET.get('desde', ''), '%Y-%m-%d').date()
    except ValueError:
        desde = hoy - timedelta(days=30)
    try:
        hasta = datetime.strptime(request.GET.get('hasta', ''), '%Y-%m-%d').date()
    except ValueError:
        hasta = hoy
    if desde > hasta:
        desde, hasta = hasta, desde
    return desde, hasta


def _pct(parte, total):
    return round(parte * 100 / total, 1) if total else 0.0


def _consulta(user, desde, hasta, status_fil=''):
    qs = Message.objects.filter(owner=user, sent_at__date__gte=desde, sent_at__date__lte=hasta)
    if status_fil:
        qs = qs.filter(delivery_status=status_fil)
    return qs


def datos_reporte(user, desde, hasta, status_fil=''):
    """Todas las cifras del reporte en un diccionario (lo usan la página, el Excel y el PDF)."""
    qs = _consulta(user, desde, hasta, status_fil)
    out = qs.filter(direction='outbound')

    conteo = {k: 0 for k, _ in ESTADOS}
    for row in out.values('delivery_status').annotate(n=Count('id')):
        conteo[row['delivery_status'] or 'sent'] = conteo.get(row['delivery_status'] or 'sent', 0) + row['n']

    enviados = sum(conteo.values())
    recibidos = qs.filter(direction='inbound').count()
    # Un mensaje leído también fue entregado
    entregados = conteo['delivered'] + conteo['read']
    leidos = conteo['read']
    fallidos = conteo['failed']

    resumen = {
        'enviados': enviados,
        'recibidos': recibidos,
        'entregados': entregados,
        'leidos': leidos,
        'fallidos': fallidos,
        'pendientes': conteo['pending'],
        'enviados_puro': conteo['sent'],
        'tasa_entrega': _pct(entregados, enviados),
        'tasa_lectura': _pct(leidos, enviados),
        'tasa_lectura_entregados': _pct(leidos, entregados),
        'tasa_fallo': _pct(fallidos, enviados),
    }

    por_estado = [
        {'estado': k, 'etiqueta': label, 'total': conteo[k], 'pct': _pct(conteo[k], enviados)}
        for k, label in ESTADOS
    ]

    # Por día (en hora de Ecuador), con el desglose por estado. Se incluyen
    # los días sin mensajes para que el gráfico no salte fechas.
    filtros_estado = {k: Count('id', filter=Q(direction='outbound', delivery_status=k)) for k, _ in ESTADOS}
    filas_dia = {
        row['dia']: row
        for row in qs.annotate(dia=TruncDate('sent_at')).values('dia').annotate(
            enviados=Count('id', filter=Q(direction='outbound')),
            recibidos=Count('id', filter=Q(direction='inbound')),
            **filtros_estado,
        ) if row['dia']
    }
    por_dia = []
    dias = (hasta - desde).days + 1
    if dias <= 120:
        rango = [desde + timedelta(days=i) for i in range(dias)]
    else:
        rango = sorted(filas_dia)
    for d in rango:
        row = filas_dia.get(d, {})
        por_dia.append({
            'dia': d.isoformat(),
            'enviados': row.get('enviados', 0),
            'recibidos': row.get('recibidos', 0),
            **{k: row.get(k, 0) for k, _ in ESTADOS},
        })

    # Por hora del día (0-23)
    horas = {
        row['hora']: row
        for row in out.annotate(hora=ExtractHour('sent_at')).values('hora').annotate(
            enviados=Count('id'),
            leidos=Count('id', filter=Q(delivery_status='read')),
            entregados=Count('id', filter=Q(delivery_status__in=['delivered', 'read'])),
        )
    }
    por_hora = []
    for h in range(24):
        row = horas.get(h, {})
        env = row.get('enviados', 0)
        por_hora.append({
            'hora': h,
            'enviados': env,
            'entregados': row.get('entregados', 0),
            'leidos': row.get('leidos', 0),
            'tasa_lectura': _pct(row.get('leidos', 0), env),
        })

    # Por campaña (los mensajes sueltos del chat van como "Mensajes directos")
    por_campana = []
    for row in out.values('campaign_id', 'campaign__name').annotate(
        enviados=Count('id'),
        entregados=Count('id', filter=Q(delivery_status__in=['delivered', 'read'])),
        leidos=Count('id', filter=Q(delivery_status='read')),
        fallidos=Count('id', filter=Q(delivery_status='failed')),
    ).order_by('-enviados'):
        env = row['enviados']
        por_campana.append({
            'campana': row['campaign__name'] or 'Mensajes directos',
            'enviados': env,
            'entregados': row['entregados'],
            'leidos': row['leidos'],
            'fallidos': row['fallidos'],
            'tasa_entrega': _pct(row['entregados'], env),
            'tasa_lectura': _pct(row['leidos'], env),
        })

    return {
        'desde': desde.isoformat(),
        'hasta': hasta.isoformat(),
        'resumen': resumen,
        'por_estado': por_estado,
        'por_dia': por_dia,
        'por_hora': por_hora,
        'por_campana': por_campana,
    }


def mensajes_detalle(user, desde, hasta, status_fil='', limite=100):
    """Mensajes salientes más recientes con el nombre del contacto (una sola consulta de contactos)."""
    qs = _consulta(user, desde, hasta, status_fil).filter(direction='outbound') \
        .select_related('campaign').order_by('-sent_at')[:limite]
    msgs = list(qs)
    telefonos = {m.phone_number for m in msgs}
    nombres = dict(
        Contact.objects.filter(owner=user, phone_number__in=telefonos).values_list('phone_number', 'full_name')
    )
    filas = []
    for m in msgs:
        local = timezone.localtime(m.sent_at) if m.sent_at else None
        filas.append({
            'id': m.id,
            'telefono': m.phone_number,
            'nombre': nombres.get(m.phone_number) or m.phone_number,
            'campana': m.campaign.name if m.campaign else '',
            'mensaje': m.content or '',
            'status': m.delivery_status or 'sent',
            'fecha': local.strftime('%d/%m/%Y %H:%M') if local else '',
            'fecha_dt': local.replace(tzinfo=None) if local else None,
        })
    return filas


def _telefono_legible(numero):
    n = normalizar_o_igual(numero)
    if len(n) == 12 and n.startswith('593'):
        return f'+593 {n[3:5]} {n[5:8]} {n[8:]}'
    return numero


def _fecha_es(iso):
    return datetime.strptime(iso, '%Y-%m-%d').strftime('%d/%m/%Y')


# ════════════════════════════════════════════════════════════════════════
# EXCEL
# ════════════════════════════════════════════════════════════════════════
def excel_reporte(user, desde, hasta, status_fil=''):
    """Libro de Excel con resumen, gráficos nativos de Excel y el detalle. Devuelve bytes."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, DoughnutChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.series import DataPoint
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    datos = datos_reporte(user, desde, hasta, status_fil)
    r = datos['resumen']

    wb = Workbook()
    titulo = Font(bold=True, size=14, color='0B0B0B')
    cab = Font(bold=True, color='FFFFFF')
    cab_fill = PatternFill('solid', fgColor='1C5CAB')
    gris = Font(color='52514E', italic=True)

    def encabezado(ws, fila, textos):
        for i, t in enumerate(textos, start=1):
            c = ws.cell(row=fila, column=i, value=t)
            c.font = cab
            c.fill = cab_fill
            c.alignment = Alignment(horizontal='center')

    def anchos(ws, valores):
        for i, w in enumerate(valores, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ── Hoja 1: Resumen ──────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Resumen'
    ws['A1'] = 'Reporte de mensajería - MassSend'
    ws['A1'].font = titulo
    ws['A2'] = f"Periodo: {_fecha_es(datos['desde'])} al {_fecha_es(datos['hasta'])}"
    ws['A3'] = f"Generado: {timezone.localtime().strftime('%d/%m/%Y %H:%M')} · Usuario: {user.get_username()}"
    ws['A2'].font = gris
    ws['A3'].font = gris

    encabezado(ws, 5, ['Indicador', 'Valor', '% de enviados'])
    kpis = [
        ('Mensajes enviados', r['enviados'], None),
        ('Entregados (incluye leídos)', r['entregados'], r['tasa_entrega'] / 100),
        ('Leídos', r['leidos'], r['tasa_lectura'] / 100),
        ('Fallidos', r['fallidos'], r['tasa_fallo'] / 100),
        ('Pendientes', r['pendientes'], _pct(r['pendientes'], r['enviados']) / 100),
        ('Mensajes recibidos', r['recibidos'], None),
        ('Leídos sobre entregados', None, r['tasa_lectura_entregados'] / 100),
    ]
    for i, (nombre, valor, pct) in enumerate(kpis, start=6):
        ws.cell(row=i, column=1, value=nombre)
        ws.cell(row=i, column=2, value=valor)
        c = ws.cell(row=i, column=3, value=pct)
        c.number_format = '0.0%'

    # Distribución por estado (para el gráfico de anillo)
    fila0 = 15
    ws.cell(row=fila0 - 1, column=1, value='Distribución por estado').font = Font(bold=True, size=12)
    encabezado(ws, fila0, ['Estado', 'Mensajes', '%'])
    for i, e in enumerate(datos['por_estado'], start=fila0 + 1):
        ws.cell(row=i, column=1, value=e['etiqueta'])
        ws.cell(row=i, column=2, value=e['total'])
        ws.cell(row=i, column=3, value=e['pct'] / 100).number_format = '0.0%'
    fin_estado = fila0 + len(datos['por_estado'])

    dona = DoughnutChart(holeSize=55)
    dona.title = 'Mensajes por estado'
    dona.add_data(Reference(ws, min_col=2, min_row=fila0, max_row=fin_estado), titles_from_data=True)
    dona.set_categories(Reference(ws, min_col=1, min_row=fila0 + 1, max_row=fin_estado))
    serie = dona.series[0]
    for idx, e in enumerate(datos['por_estado']):
        pt = DataPoint(idx=idx)
        pt.graphicalProperties.solidFill = COLOR[e['estado']][1:]
        pt.graphicalProperties.line.solidFill = 'FFFFFF'
        serie.dPt.append(pt)
    dona.dataLabels = DataLabelList()
    dona.dataLabels.showPercent = True
    dona.dataLabels.showVal = False
    dona.dataLabels.showCatName = False
    dona.dataLabels.showSerName = False
    dona.dataLabels.showLegendKey = False
    dona.legend.position = 'r'
    dona.height, dona.width = 8, 12
    ws.add_chart(dona, 'E5')

    # Embudo: enviados -> entregados -> leídos
    fila_e = fin_estado + 3
    ws.cell(row=fila_e - 1, column=1, value='Embudo de entrega').font = Font(bold=True, size=12)
    encabezado(ws, fila_e, ['Etapa', 'Mensajes', '% de enviados'])
    etapas = [('Enviados', r['enviados'], 1 if r['enviados'] else 0),
              ('Entregados', r['entregados'], r['tasa_entrega'] / 100),
              ('Leídos', r['leidos'], r['tasa_lectura'] / 100)]
    for i, (n, v, p) in enumerate(etapas, start=fila_e + 1):
        ws.cell(row=i, column=1, value=n)
        ws.cell(row=i, column=2, value=v)
        ws.cell(row=i, column=3, value=p).number_format = '0.0%'
    embudo = BarChart()
    embudo.type = 'bar'
    embudo.title = 'Embudo de entrega'
    embudo.add_data(Reference(ws, min_col=2, min_row=fila_e, max_row=fila_e + 3), titles_from_data=True)
    embudo.set_categories(Reference(ws, min_col=1, min_row=fila_e + 1, max_row=fila_e + 3))
    embudo.y_axis.majorGridlines = None
    embudo.x_axis.scaling.orientation = 'maxMin'   # Enviados arriba
    embudo.legend = None
    embudo.gapWidth = 40
    s = embudo.series[0]
    for idx, color in enumerate(['86B6EF', '2A78D6', '104281']):
        pt = DataPoint(idx=idx)
        pt.graphicalProperties.solidFill = color
        s.dPt.append(pt)
    embudo.dataLabels = DataLabelList()
    embudo.dataLabels.showVal = True
    embudo.dataLabels.showCatName = False
    embudo.dataLabels.showSerName = False
    embudo.dataLabels.showLegendKey = False
    embudo.y_axis.scaling.min = 0
    embudo.height, embudo.width = 6, 12
    ws.add_chart(embudo, 'E22')
    anchos(ws, [30, 12, 14])

    # ── Hoja 2: Por día ──────────────────────────────────────────────
    ws = wb.create_sheet('Por día')
    cols = ['Fecha'] + [label for _, label in ESTADOS] + ['Total enviados', 'Recibidos']
    encabezado(ws, 1, cols)
    for i, d in enumerate(datos['por_dia'], start=2):
        ws.cell(row=i, column=1, value=datetime.strptime(d['dia'], '%Y-%m-%d').date()).number_format = 'DD/MM/YYYY'
        for j, (k, _) in enumerate(ESTADOS, start=2):
            ws.cell(row=i, column=j, value=d[k])
        ws.cell(row=i, column=len(ESTADOS) + 2, value=d['enviados'])
        ws.cell(row=i, column=len(ESTADOS) + 3, value=d['recibidos'])
    ws.freeze_panes = 'A2'
    anchos(ws, [12] + [12] * len(ESTADOS) + [15, 11])
    n_dias = len(datos['por_dia'])
    if n_dias:
        barras = BarChart()
        barras.type = 'col'
        barras.grouping = 'stacked'
        barras.overlap = 100
        barras.title = 'Mensajes enviados por día y estado'
        barras.add_data(Reference(ws, min_col=2, max_col=len(ESTADOS) + 1, min_row=1, max_row=n_dias + 1),
                        titles_from_data=True)
        barras.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n_dias + 1))
        for s, (k, _) in zip(barras.series, ESTADOS):
            s.graphicalProperties.solidFill = COLOR[k][1:]
            s.graphicalProperties.line.noFill = True
        barras.gapWidth = 50
        barras.y_axis.scaling.min = 0
        barras.height, barras.width = 9, 22
        ws.add_chart(barras, f'{get_column_letter(len(cols) + 2)}2')

    # ── Hoja 3: Por hora ─────────────────────────────────────────────
    ws = wb.create_sheet('Por hora')
    encabezado(ws, 1, ['Hora', 'Enviados', 'Entregados', 'Leídos', 'Tasa de lectura'])
    for i, h in enumerate(datos['por_hora'], start=2):
        ws.cell(row=i, column=1, value=f"{h['hora']:02d}:00")
        ws.cell(row=i, column=2, value=h['enviados'])
        ws.cell(row=i, column=3, value=h['entregados'])
        ws.cell(row=i, column=4, value=h['leidos'])
        ws.cell(row=i, column=5, value=h['tasa_lectura'] / 100).number_format = '0.0%'
    anchos(ws, [9, 11, 12, 9, 15])
    hora = BarChart()
    hora.type = 'col'
    hora.title = 'Envíos por hora del día'
    hora.add_data(Reference(ws, min_col=2, min_row=1, max_row=25), titles_from_data=True)
    hora.set_categories(Reference(ws, min_col=1, min_row=2, max_row=25))
    hora.series[0].graphicalProperties.solidFill = '2A78D6'
    hora.legend = None
    hora.y_axis.scaling.min = 0
    hora.gapWidth = 40
    hora.height, hora.width = 8, 20
    ws.add_chart(hora, 'G2')

    # ── Hoja 4: Campañas ─────────────────────────────────────────────
    ws = wb.create_sheet('Campañas')
    encabezado(ws, 1, ['Campaña', 'Enviados', 'Entregados', 'Leídos', 'Fallidos', 'Tasa de entrega', 'Tasa de lectura'])
    for i, c in enumerate(datos['por_campana'], start=2):
        ws.cell(row=i, column=1, value=c['campana'])
        ws.cell(row=i, column=2, value=c['enviados'])
        ws.cell(row=i, column=3, value=c['entregados'])
        ws.cell(row=i, column=4, value=c['leidos'])
        ws.cell(row=i, column=5, value=c['fallidos'])
        ws.cell(row=i, column=6, value=c['tasa_entrega'] / 100).number_format = '0.0%'
        ws.cell(row=i, column=7, value=c['tasa_lectura'] / 100).number_format = '0.0%'
    ws.freeze_panes = 'A2'
    anchos(ws, [32, 11, 12, 9, 10, 15, 15])

    # ── Hoja 5: Mensajes (detalle) ───────────────────────────────────
    ws = wb.create_sheet('Mensajes')
    encabezado(ws, 1, ['Fecha', 'Contacto', 'Teléfono', 'Campaña', 'Estado', 'Mensaje'])
    filas = mensajes_detalle(user, desde, hasta, status_fil, limite=MAX_FILAS_EXCEL)
    for i, m in enumerate(filas, start=2):
        ws.cell(row=i, column=1, value=m['fecha_dt']).number_format = 'DD/MM/YYYY HH:MM'
        ws.cell(row=i, column=2, value=m['nombre'])
        ws.cell(row=i, column=3, value=_telefono_legible(m['telefono']))
        ws.cell(row=i, column=4, value=m['campana'] or 'Mensaje directo')
        ws.cell(row=i, column=5, value=ETIQUETA.get(m['status'], m['status']))
        ws.cell(row=i, column=6, value=m['mensaje'][:500])
    ws.freeze_panes = 'A2'
    if filas:
        ws.auto_filter.ref = f'A1:F{len(filas) + 1}'
    if len(filas) >= MAX_FILAS_EXCEL:
        ws.cell(row=len(filas) + 3, column=1,
                value=f'Se muestran los {MAX_FILAS_EXCEL} mensajes más recientes del periodo.').font = gris
    anchos(ws, [17, 24, 18, 24, 12, 70])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ════════════════════════════════════════════════════════════════════════
# PDF
# ════════════════════════════════════════════════════════════════════════
def pdf_reporte(user, desde, hasta, status_fil=''):
    """Informe en PDF (A4) con indicadores, gráficos y tablas. Devuelve bytes."""
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
    from reportlab.graphics.charts.doughnut import Doughnut
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    datos = datos_reporte(user, desde, hasta, status_fil)
    r = datos['resumen']
    hx = colors.HexColor
    TINTA, TINTA2, MUTED, RAYA, FONDO = hx('#0b0b0b'), hx('#52514e'), hx('#8a8983'), hx('#e4e3df'), hx('#f6f6f4')

    estilos = getSampleStyleSheet()
    h1 = ParagraphStyle('h1', parent=estilos['Title'], alignment=TA_LEFT, fontSize=20, leading=24,
                        textColor=TINTA, spaceAfter=2)
    h2 = ParagraphStyle('h2', parent=estilos['Heading2'], fontSize=12.5, leading=16, textColor=TINTA,
                        spaceBefore=14, spaceAfter=6)
    sub = ParagraphStyle('sub', parent=estilos['Normal'], fontSize=9, textColor=TINTA2, leading=12)
    celda = ParagraphStyle('celda', parent=estilos['Normal'], fontSize=7.8, leading=9.5, textColor=TINTA)
    nota = ParagraphStyle('nota', parent=sub, fontSize=8, textColor=MUTED)
    cifra = ParagraphStyle('cifra', parent=estilos['Normal'], fontName='Helvetica-Bold', fontSize=18,
                           leading=22, textColor=TINTA)

    ancho = A4[0] - 3.2 * cm
    story = [
        Paragraph('Reporte de mensajería', h1),
        Paragraph(f"Periodo {_fecha_es(datos['desde'])} al {_fecha_es(datos['hasta'])} · "
                  f"Generado el {timezone.localtime().strftime('%d/%m/%Y %H:%M')} · "
                  f"{user.get_username()}", sub),
        Spacer(1, 12),
    ]

    # ── Indicadores ──────────────────────────────────────────────────
    def tile(valor, etiqueta, detalle):
        return [Paragraph(str(valor), cifra),
                Paragraph(etiqueta, sub),
                Paragraph(detalle, nota)]
    tiles = [
        tile(r['enviados'], 'Enviados', f"{r['recibidos']} recibidos"),
        tile(r['entregados'], 'Entregados', f"{r['tasa_entrega']}% de enviados"),
        tile(r['leidos'], 'Leídos', f"{r['tasa_lectura']}% de enviados"),
        tile(r['fallidos'], 'Fallidos', f"{r['tasa_fallo']}% de enviados"),
    ]
    t = Table([tiles], colWidths=[ancho / 4] * 4)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), FONDO),
        ('LINEAFTER', (0, 0), (-2, -1), 2, colors.white),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
    ]))
    story.append(t)

    # ── Distribución por estado (anillo) + embudo ───────────────────
    estados_con_datos = [e for e in datos['por_estado'] if e['total'] > 0]
    d = Drawing(ancho, 170)
    if estados_con_datos:
        dona = Doughnut()
        dona.x, dona.y, dona.width, dona.height = 0, 10, 150, 150
        dona.data = [e['total'] for e in estados_con_datos]
        dona.labels = None
        dona.innerRadiusFraction = 0.58
        dona.strokeColor = colors.white
        dona.strokeWidth = 2
        for i, e in enumerate(estados_con_datos):
            dona.slices[i].fillColor = hx(COLOR[e['estado']])
            dona.slices[i].strokeColor = colors.white
            dona.slices[i].strokeWidth = 2
        d.add(dona)
        d.add(String(75, 88, str(r['enviados']), fontName='Helvetica-Bold', fontSize=16,
                     fillColor=TINTA, textAnchor='middle'))
        d.add(String(75, 74, 'enviados', fontName='Helvetica', fontSize=8, fillColor=TINTA2,
                     textAnchor='middle'))
        # Leyenda con cifras (la identidad nunca va solo por color)
        y = 140
        for e in datos['por_estado']:
            d.add(Rect(165, y, 8, 8, fillColor=hx(COLOR[e['estado']]), strokeColor=None))
            d.add(String(178, y + 1, e['etiqueta'], fontName='Helvetica', fontSize=8.5, fillColor=TINTA))
            d.add(String(258, y + 1, f"{e['total']}  ({e['pct']}%)", fontName='Helvetica', fontSize=8.5,
                         fillColor=TINTA2))
            y -= 16

        # Embudo
        emb = HorizontalBarChart()
        emb.x, emb.y, emb.width, emb.height = 390, 30, ancho - 400, 110
        emb.data = [[r['leidos'], r['entregados'], r['enviados']]]
        emb.categoryAxis.categoryNames = ['Leídos', 'Entregados', 'Enviados']
        emb.categoryAxis.labels.fontSize = 8
        emb.categoryAxis.labels.fontName = 'Helvetica'
        emb.categoryAxis.labels.fillColor = TINTA2
        emb.categoryAxis.strokeColor = RAYA
        emb.valueAxis.valueMin = 0
        emb.valueAxis.visible = False
        emb.bars.strokeColor = None
        emb.barSpacing = 6
        for i, c in enumerate(['#104281', '#2a78d6', '#86b6ef']):
            emb.bars[(0, i)].fillColor = hx(c)
        emb.barLabelFormat = '%d'
        emb.barLabels.fontSize = 8
        emb.barLabels.fontName = 'Helvetica'
        emb.barLabels.fillColor = TINTA
        emb.barLabels.boxAnchor = 'w'
        emb.barLabels.dx = 4
        d.add(emb)
        d.add(String(390, 152, 'Embudo de entrega', fontName='Helvetica-Bold', fontSize=9, fillColor=TINTA))
        story += [Paragraph('Mensajes por estado', h2), d]
    else:
        story += [Spacer(1, 16), Paragraph('No hay mensajes enviados en este periodo.', sub)]

    # ── Por día (columnas apiladas por estado) ──────────────────────
    dias = datos['por_dia']
    if any(x['enviados'] for x in dias):
        alto = 190
        g = Drawing(ancho, alto + 20)
        bc = VerticalBarChart()
        bc.x, bc.y, bc.width, bc.height = 30, 30, ancho - 40, alto - 40
        bc.data = [[x[k] for x in dias] for k, _ in ESTADOS]
        bc.categoryAxis.style = 'stacked'
        cada = max(1, len(dias) // 10)
        bc.categoryAxis.categoryNames = [
            datetime.strptime(x['dia'], '%Y-%m-%d').strftime('%d/%m') if i % cada == 0 else ''
            for i, x in enumerate(dias)
        ]
        bc.categoryAxis.labels.fontSize = 7
        bc.categoryAxis.labels.fontName = 'Helvetica'
        bc.categoryAxis.labels.fillColor = TINTA2
        bc.categoryAxis.strokeColor = RAYA
        bc.categoryAxis.tickDown = 0
        bc.valueAxis.valueMin = 0
        bc.valueAxis.labels.fontSize = 7
        bc.valueAxis.labels.fontName = 'Helvetica'
        bc.valueAxis.labels.fillColor = TINTA2
        bc.valueAxis.strokeColor = None
        bc.valueAxis.gridStrokeColor = RAYA
        bc.valueAxis.visibleGrid = True
        bc.valueAxis.forceZero = True
        maximo = max(x['enviados'] for x in dias)
        bc.valueAxis.valueStep = max(1, round(maximo / 4)) if maximo > 4 else 1
        bc.bars.strokeColor = colors.white
        bc.bars.strokeWidth = 0.6
        bc.groupSpacing = 3 if len(dias) > 20 else 8
        for i, (k, _) in enumerate(ESTADOS):
            bc.bars[i].fillColor = hx(COLOR[k])
        g.add(bc)
        x = 30
        for k, label in ESTADOS:
            g.add(Rect(x, alto + 4, 8, 8, fillColor=hx(COLOR[k]), strokeColor=None))
            g.add(String(x + 12, alto + 5, label, fontName='Helvetica', fontSize=8, fillColor=TINTA2))
            x += 78
        story += [Paragraph('Actividad diaria (mensajes enviados por estado)', h2), g]

    # ── Por hora del día ────────────────────────────────────────────
    horas = datos['por_hora']
    if any(h['enviados'] for h in horas):
        g = Drawing(ancho, 160)
        hb = VerticalBarChart()
        hb.x, hb.y, hb.width, hb.height = 30, 25, ancho - 40, 120
        hb.data = [[h['enviados'] for h in horas]]
        hb.categoryAxis.categoryNames = [f"{h['hora']:02d}" if h['hora'] % 2 == 0 else '' for h in horas]
        hb.categoryAxis.labels.fontSize = 7
        hb.categoryAxis.labels.fontName = 'Helvetica'
        hb.categoryAxis.labels.fillColor = TINTA2
        hb.categoryAxis.strokeColor = RAYA
        hb.categoryAxis.tickDown = 0
        hb.valueAxis.valueMin = 0
        hb.valueAxis.labels.fontSize = 7
        hb.valueAxis.labels.fontName = 'Helvetica'
        hb.valueAxis.labels.fillColor = TINTA2
        hb.valueAxis.strokeColor = None
        hb.valueAxis.visibleGrid = True
        hb.valueAxis.gridStrokeColor = RAYA
        maximo = max(h['enviados'] for h in horas)
        hb.valueAxis.valueStep = max(1, round(maximo / 4)) if maximo > 4 else 1
        hb.bars[0].fillColor = hx('#2a78d6')
        hb.bars.strokeColor = None
        hb.groupSpacing = 4
        g.add(hb)
        mejor = max(horas, key=lambda h: (h['enviados'] and h['tasa_lectura'], h['enviados']))
        pie = (f"Hora con mejor tasa de lectura: {mejor['hora']:02d}:00 "
               f"({mejor['tasa_lectura']}% de {mejor['enviados']} enviados)") if mejor['enviados'] else ''
        story += [KeepTogether([Paragraph('Envíos por hora del día', h2), g, Paragraph(pie, nota)])]

    # ── Campañas ────────────────────────────────────────────────────
    camp = datos['por_campana']
    if camp:
        filas = [['Campaña', 'Enviados', 'Entregados', 'Leídos', 'Fallidos', '% entrega', '% lectura']]
        for c in camp:
            filas.append([Paragraph(c['campana'], celda), c['enviados'], c['entregados'], c['leidos'],
                          c['fallidos'], f"{c['tasa_entrega']}%", f"{c['tasa_lectura']}%"])
        t = Table(filas, colWidths=[ancho * 0.34] + [ancho * 0.11] * 6, repeatRows=1)
        t.setStyle(_estilo_tabla(colors, TINTA, TINTA2, RAYA, FONDO, numericas_desde=1))
        story += [Paragraph('Rendimiento por campaña', h2), t]

    # ── Detalle de mensajes ─────────────────────────────────────────
    msgs = mensajes_detalle(user, desde, hasta, status_fil, limite=MAX_FILAS_PDF)
    if msgs:
        filas = [['Fecha', 'Contacto', 'Teléfono', 'Campaña', 'Estado', 'Mensaje']]
        for m in msgs:
            texto = m['mensaje'].replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            filas.append([
                m['fecha'], Paragraph(m['nombre'].replace('&', '&amp;').replace('<', '&lt;'), celda),
                _telefono_legible(m['telefono']),
                Paragraph((m['campana'] or 'Directo').replace('&', '&amp;').replace('<', '&lt;'), celda),
                ETIQUETA.get(m['status'], m['status']),
                Paragraph(texto[:90] + ('…' if len(texto) > 90 else ''), celda),
            ])
        t = Table(filas, colWidths=[ancho * w for w in (0.15, 0.16, 0.16, 0.15, 0.1, 0.28)], repeatRows=1)
        t.setStyle(_estilo_tabla(colors, TINTA, TINTA2, RAYA, FONDO))
        titulo = 'Detalle de mensajes'
        if len(msgs) >= MAX_FILAS_PDF:
            titulo += f' (los {MAX_FILAS_PDF} más recientes; el detalle completo está en el Excel)'
        story += [Paragraph(titulo, h2), t]

    def pie_pagina(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 7.5)
        canvas.setFillColor(MUTED)
        canvas.drawString(1.6 * cm, 1.0 * cm, 'MassSend · Reporte de mensajería')
        canvas.drawRightString(A4[0] - 1.6 * cm, 1.0 * cm, f'Página {doc.page}')
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.6 * cm,
                            title='Reporte de mensajería', author='MassSend')
    doc.build(story, onFirstPage=pie_pagina, onLaterPages=pie_pagina)
    return buf.getvalue()


def _estilo_tabla(colors, TINTA, TINTA2, RAYA, FONDO, numericas_desde=None):
    from reportlab.platypus import TableStyle
    estilo = [
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7.8),
        ('TEXTCOLOR', (0, 0), (-1, 0), TINTA2),
        ('TEXTCOLOR', (0, 1), (-1, -1), TINTA),
        ('LINEBELOW', (0, 0), (-1, 0), 0.8, TINTA2),
        ('LINEBELOW', (0, 1), (-1, -1), 0.3, RAYA),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, FONDO]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]
    if numericas_desde is not None:
        estilo.append(('ALIGN', (numericas_desde, 0), (-1, -1), 'RIGHT'))
    return TableStyle(estilo)
